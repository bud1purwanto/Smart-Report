import pytest
import pandas as pd
from app.services.pandas_engine import pandas_engine
from app.services.pandas_engine import DuplicateComparisonKeyError
from app.services.abap_validator import abap_validator
from app.schemas.query import QueryDefinition, TableNodeItem, JoinItem, FieldSelectionItem, FilterItem
from app.core.security import encrypt_password, decrypt_password

def test_encryption_decryption():
    raw = "super_secret_sap_pass_123"
    enc = encrypt_password(raw)
    assert enc != raw
    dec = decrypt_password(enc)
    assert dec == raw

def test_pandas_formulas():
    df = pd.DataFrame([
        {"NETPR": 100, "MENGE": 5},
        {"NETPR": 250, "MENGE": 2}
    ])
    formulas = [{"name": "TOTAL", "formula": "row['NETPR'] * row['MENGE']"}]
    res = pandas_engine.apply_custom_formulas(df, formulas)
    assert "TOTAL" in res.columns
    assert res["TOTAL"].tolist() == [500, 500]

def test_pandas_deduplication():
    df = pd.DataFrame([
        {"EBELN": "4500000001", "EBELP": "00010", "VAL": 10},
        {"EBELN": "4500000001", "EBELP": "00010", "VAL": 10},
        {"EBELN": "4500000002", "EBELP": "00010", "VAL": 20}
    ])
    deduped = pandas_engine.deduplicate(df, subset_keys=["EBELN", "EBELP"])
    assert len(deduped) == 2

def test_pandas_anonymization():
    df = pd.DataFrame([
        {"LIFNR": "1000", "NAME1": "Vendor A", "BANKN": "1234567890", "WRBTR": 5000000.00}
    ])
    masked = pandas_engine.anonymize(df)
    assert masked["BANKN"].iloc[0] == "***7890"
    assert masked["WRBTR"].iloc[0] == "***.***,00"
    assert masked["NAME1"].iloc[0] == "Vendor A"

def test_pandas_diff_datasets():
    df_a = pd.DataFrame([
        {"MATNR": "MAT-01", "MAKTX": "Widget Standard", "PRICE": 100},
        {"MATNR": "MAT-02", "MAKTX": "Bolt 10mm", "PRICE": 50},
        {"MATNR": "MAT-03", "MAKTX": "Nut 10mm", "PRICE": 25}
    ])
    df_b = pd.DataFrame([
        {"MATNR": "MAT-01", "MAKTX": "Widget Standard", "PRICE": 100}, # Identical
        {"MATNR": "MAT-02", "MAKTX": "Bolt 10mm - Heavy", "PRICE": 60}, # Modified
        {"MATNR": "MAT-04", "MAKTX": "Washer 10mm", "PRICE": 15} # Added in B, MAT-03 Deleted
    ])
    summary, diffs = pandas_engine.diff_datasets(df_a, df_b, key_fields=["MATNR"])
    assert summary["identical_count"] == 1
    assert summary["modified_count"] == 1
    assert summary["added_count"] == 1
    assert summary["deleted_count"] == 1


def test_pandas_diff_rejects_duplicate_composite_keys():
    df_a = pd.DataFrame([
        {"MATNR": "MAT-01", "PRICE": 100},
        {"MATNR": "MAT-01", "PRICE": 110},
    ])
    df_b = pd.DataFrame([{"MATNR": "MAT-01", "PRICE": 100}])

    with pytest.raises(DuplicateComparisonKeyError):
        pandas_engine.diff_datasets(df_a, df_b, key_fields=["MATNR"])

def test_excel_export():
    df = pd.DataFrame([
        {"MATNR": "MAT-01", "WRBTR": 1000}
    ])
    content = pandas_engine.export_to_excel(df, title="Test", anonymize=True)
    assert isinstance(content, bytes)
    assert len(content) > 0


def test_excel_export_always_masks_sensitive_financial_fields():
    import io
    from openpyxl import load_workbook

    df = pd.DataFrame([{"MATNR": "MAT-01", "WRBTR": 1250000, "BANKN": "1234567890"}])
    content = pandas_engine.export_to_excel(df, title="Sensitive", anonymize=False)
    sheet = load_workbook(io.BytesIO(content), read_only=True).active
    values = list(sheet.iter_rows(values_only=True))

    assert values[1][1] == "***.***,00"
    assert values[1][2] == "***7890"

def test_abap_validator_valid():
    query = QueryDefinition(
        tables=[
            TableNodeItem(id="t1", table="EKKO"),
            TableNodeItem(id="t2", table="EKPO")
        ],
        joins=[
            JoinItem(id="j1", sourceTableId="t1", targetTableId="t2", sourceField="EBELN", targetField="EBELN", joinType="INNER")
        ],
        selectedFields=[
            FieldSelectionItem(tableId="t1", table="EKKO", field="EBELN"),
            FieldSelectionItem(tableId="t2", table="EKPO", field="EBELP")
        ],
        filters=[
            FilterItem(field="EKKO.BSTYP", operator="EQ", value="F")
        ]
    )
    val = abap_validator.validate_and_generate_sql(query)
    assert val["is_valid"] is True
    assert "INNER JOIN ekpo" in val["open_sql"]
    assert "WHERE ekko~bstyp = 'F'." in val["open_sql"]

def test_abap_validator_cartesian():
    # 2 tables without join
    query = QueryDefinition(
        tables=[
            TableNodeItem(id="t1", table="EKKO"),
            TableNodeItem(id="t2", table="MARA")
        ],
        joins=[],
        selectedFields=[]
    )
    val = abap_validator.validate_and_generate_sql(query)
    assert val["is_valid"] is False
    assert any("Cartesian Product" in err for err in val["errors"])


def test_query_executor_build_rfc_where_clauses():
    from app.services.query_executor import build_rfc_where_clauses, parse_filter_condition

    # Test short clauses
    clauses = ["EBELN = '4500000001'", "BUKRS = '1000'"]
    options = build_rfc_where_clauses(clauses, connector="AND")
    assert len(options) == 2
    assert options[0].strip().endswith("AND")
    assert "BUKRS = '1000'" in options[1]
    for opt in options:
        assert len(opt) <= 72

    # Test long clauses wrapping
    long_clauses = [f"EBELN = '45000000{i:02d}'" for i in range(25)]
    wrapped_options = build_rfc_where_clauses(long_clauses, connector="OR")
    for opt in wrapped_options:
        assert len(opt) <= 72
        assert not opt.startswith("AND ")
        assert not opt.startswith("OR ")

    # Test parse_filter_condition
    flt = FilterItem(field="EKKO.BSTYP", operator="EQ", value="F")
    cond = parse_filter_condition(flt, "EKKO")
    assert cond == "BSTYP = 'F'"

    flt_between = FilterItem(field="BSART", operator="BETWEEN", value="NB", valueTo="UB")
    cond_between = parse_filter_condition(flt_between, "EKKO")
    assert cond_between == "BSART BETWEEN 'NB' AND 'UB'"

    flt_in = FilterItem(field="EKKO.BSART", operator="IN", value="NB, UB, FO")
    cond_in = parse_filter_condition(flt_in, "EKKO")
    assert cond_in == "BSART IN ('NB', 'UB', 'FO')"

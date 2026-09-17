import io
import re
import math
from typing import List, Dict, Any, Optional, Tuple
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Sensitive financial, tax, and banking fields for vendor data masking
SENSITIVE_VENDOR_FIELDS = {
    "BANKN", "IBAN", "SWIFT", "STCD1", "STCD2", "STCD3", "STCD4",
    "WRBTR", "DMBTR", "PSWBT", "NETPR", "NETWR", "BRTWR", "KBETR",
    "WAERS_VAL", "SALES_VAL", "AMOUNT", "SALARY", "PRICE"
}

class PandasEngine:
    """
    Pandas-based high-performance in-memory calculation, diffing, deduplication, 
    anonymization, and styled Excel generation engine.
    """

    @staticmethod
    def deduplicate(df: pd.DataFrame, subset_keys: Optional[List[str]] = None) -> pd.DataFrame:
        """Removes duplicate rows using Pandas drop_duplicates."""
        if df.empty:
            return df
        if subset_keys:
            valid_keys = [k for k in subset_keys if k in df.columns]
            if valid_keys:
                return df.drop_duplicates(subset=valid_keys, keep="first")
        return df.drop_duplicates(keep="first")

    @staticmethod
    def mask_financial_value(val: Any, col_name: str = "") -> str:
        """Masks financial nominals and vendor banking numbers."""
        if pd.isna(val) or val is None or val == "":
            return ""
        s = str(val).strip()
        col_upper = col_name.upper()

        # Bank, Tax, Account IDs: mask preserving last 4 digits
        if any(term in col_upper for term in ["BANK", "IBAN", "STCD", "TAX", "ACCOUNT"]):
            if len(s) > 4:
                return f"***{s[-4:]}"
            return "****"

        # Financial nominal / currency amounts
        try:
            float(s)
            return "***.***,00"
        except ValueError:
            pass

        if len(s) > 4:
            return f"***{s[-4:]}"
        return "****"

    @classmethod
    def anonymize(
        cls,
        df: pd.DataFrame,
        target_columns: Optional[List[str]] = None
    ) -> pd.DataFrame:
        """
        Anonymizes sensitive financial and vendor data.
        RULE 4: Must run before Excel export to hide sensitive vendor financials.
        """
        if df.empty:
            return df
        
        masked_df = df.copy()
        cols_to_check = target_columns if target_columns else masked_df.columns

        for col in cols_to_check:
            col_upper = col.upper()
            if col_upper in SENSITIVE_VENDOR_FIELDS or any(sens in col_upper for sens in ["BANK", "IBAN", "TAX", "SALARY", "NOMINAL", "NETPR", "WRBTR", "DMBTR"]):
                masked_df[col] = masked_df[col].apply(lambda v: cls.mask_financial_value(v, col))

        return masked_df

    @staticmethod
    def apply_custom_formulas(
        df: pd.DataFrame,
        formulas: List[Dict[str, Any]]
    ) -> pd.DataFrame:
        """
        Calculates custom formula columns dynamically on Pandas DataFrame.
        Formulas e.g.: {"name": "TOTAL", "formula": "row.NETPR * row.MENGE"}
        or {"name": "DISCOUNT_PRICE", "formula": "row['NETPR'] * 0.9"}
        """
        if df.empty or not formulas:
            return df

        result_df = df.copy()

        for item in formulas:
            col_name = item.get("name")
            formula = item.get("formula", "")
            if not col_name or not formula:
                continue

            # Convert formula format 'row.COL' or 'row["COL"]' into df evaluation
            # Example: 'row.Quantity * row.Net_Price' -> df['Quantity'] * df['Net_Price']
            def eval_row(row):
                try:
                    # Provide local context for row evaluation
                    context = {"row": row, "math": math}
                    # Replace row.FIELD with row['FIELD'] if needed
                    expr = re.sub(r'row\.([A-Za-z0-9_]+)', r'row["\1"]', formula)
                    return eval(expr, {"__builtins__": {}}, context)
                except Exception:
                    return None

            try:
                # First try vectorized pandas eval if valid expression
                clean_expr = re.sub(r'row\[[\'"]([A-Za-z0-9_]+)[\'"]\]', r'`\1`', formula)
                clean_expr = re.sub(r'row\.([A-Za-z0-9_]+)', r'`\1`', clean_expr)
                result_df[col_name] = result_df.eval(clean_expr)
            except Exception:
                # Fallback to apply eval_row
                result_df[col_name] = result_df.apply(eval_row, axis=1)

        return result_df

    @classmethod
    def diff_datasets(
        cls,
        df_a: pd.DataFrame,
        df_b: pd.DataFrame,
        key_fields: Optional[List[str]] = None
    ) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
        """
        Cross-Server Data Compare logic:
        Performs full-outer alignment between server A and server B dataframes,
        detects IDENTICAL, MODIFIED, ADDED_IN_B, and DELETED_IN_B.
        Highlights exact field level differences.
        """
        if df_a.empty and df_b.empty:
            summary = {
                "total_a": 0, "total_b": 0,
                "identical_count": 0, "modified_count": 0,
                "added_count": 0, "deleted_count": 0
            }
            return summary, []

        # Determine key fields
        all_cols = list(set(df_a.columns).union(set(df_b.columns)))
        if not key_fields:
            # Auto-detect keys (like EBELN, MATNR, VBELN, or first column)
            candidate_keys = [c for c in all_cols if c.upper() in ["EBELN", "MATNR", "VBELN", "LIFNR", "KUNNR", "BELNR", "POSNR", "EBELP"]]
            keys = candidate_keys if candidate_keys else [all_cols[0]]
        else:
            keys = [k for k in key_fields if k in all_cols]
            if not keys and all_cols:
                keys = [all_cols[0]]

        # Ensure all cols exist in both frames
        for c in all_cols:
            if c not in df_a.columns:
                df_a[c] = None
            if c not in df_b.columns:
                df_b[c] = None

        # Create composite key representation
        def make_key_str(row):
            return "|".join([str(row.get(k, "")).strip() for k in keys])

        df_a_indexed = df_a.copy()
        df_b_indexed = df_b.copy()

        df_a_indexed["_diff_key_"] = df_a_indexed.apply(make_key_str, axis=1)
        df_b_indexed["_diff_key_"] = df_b_indexed.apply(make_key_str, axis=1)

        dict_a = {row["_diff_key_"]: row.to_dict() for _, row in df_a_indexed.iterrows()}
        dict_b = {row["_diff_key_"]: row.to_dict() for _, row in df_b_indexed.iterrows()}

        all_keys = set(dict_a.keys()).union(set(dict_b.keys()))

        identical_count = 0
        modified_count = 0
        added_count = 0
        deleted_count = 0
        diff_rows = []

        compare_cols = [c for c in all_cols if c != "_diff_key_"]

        for k in sorted(all_keys):
            in_a = k in dict_a
            in_b = k in dict_b

            if in_a and not in_b:
                deleted_count += 1
                diff_rows.append({
                    "diff_status": "DELETED_IN_B",
                    "key_value": k,
                    "data_a": dict_a[k],
                    "data_b": None,
                    "changed_fields": {}
                })
            elif not in_a and in_b:
                added_count += 1
                diff_rows.append({
                    "diff_status": "ADDED_IN_B",
                    "key_value": k,
                    "data_a": None,
                    "data_b": dict_b[k],
                    "changed_fields": {}
                })
            else:
                row_a = dict_a[k]
                row_b = dict_b[k]
                changes = {}
                for col in compare_cols:
                    val_a = str(row_a.get(col, "")).strip() if row_a.get(col) is not None else ""
                    val_b = str(row_b.get(col, "")).strip() if row_b.get(col) is not None else ""
                    if val_a != val_b:
                        changes[col] = {"old_val": row_a.get(col), "new_val": row_b.get(col)}

                if not changes:
                    identical_count += 1
                    diff_rows.append({
                        "diff_status": "IDENTICAL",
                        "key_value": k,
                        "data_a": row_a,
                        "data_b": row_b,
                        "changed_fields": {}
                    })
                else:
                    modified_count += 1
                    diff_rows.append({
                        "diff_status": "MODIFIED",
                        "key_value": k,
                        "data_a": row_a,
                        "data_b": row_b,
                        "changed_fields": changes
                    })

        summary = {
            "total_a": len(df_a),
            "total_b": len(df_b),
            "identical_count": identical_count,
            "modified_count": modified_count,
            "added_count": added_count,
            "deleted_count": deleted_count
        }

        return summary, diff_rows

    @classmethod
    def export_to_excel(
        cls,
        df: pd.DataFrame,
        title: str = "Smart_SQVI_Report",
        anonymize: bool = False,
        deduplicate: bool = False,
        dedup_keys: Optional[List[str]] = None
    ) -> bytes:
        """
        Exports DataFrame to highly polished Excel (.xlsx) file.
        Applies deduplication and sensitive financial anonymization beforehand.
        """
        work_df = df.copy()

        if deduplicate:
            work_df = cls.deduplicate(work_df, dedup_keys)

        if anonymize:
            work_df = cls.anonymize(work_df)

        wb = Workbook()
        ws = wb.active
        ws.title = title[:30] # Excel sheet title max 31 chars

        # SAP Blue ALV theme
        header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Segoe UI", size=10)
        border_thin = Side(border_style="thin", color="D3D3D3")
        cell_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)
        zebra_fill = PatternFill(start_color="F7F9FC", end_color="F7F9FC", fill_type="solid")

        # Write Header
        headers = list(work_df.columns)
        ws.append(headers)
        for col_num, header_val in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cell_border

        # Write Data
        for row_idx, row in enumerate(work_df.itertuples(index=False), 2):
            ws.append(list(row))
            is_even = (row_idx % 2 == 0)
            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.font = data_font
                cell.border = cell_border
                if is_even:
                    cell.fill = zebra_fill

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        # Enable Auto-Filter
        if headers:
            ws.auto_filter.ref = ws.dimensions

        out = io.BytesIO()
        wb.save(out)
        out.seek(0)
        return out.getvalue()

pandas_engine = PandasEngine()
